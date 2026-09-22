from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\EDA挑战题目六")
OUT = ROOT / "output" / "reports" / "PCB原理图自动解析项目综合技术报告_2026-09-16.docx"

BLUE = "1F4E78"
MID_BLUE = "5B9BD5"
PALE_BLUE = "DDEBF7"
PALEST_BLUE = "EEF4FA"
LIGHT_GRAY = "D9E2F3"
TEXT = "1F1F1F"
MUTED = "666666"
GREEN = "2E7D32"
AMBER = "9C6500"
RED = "C00000"
WHITE = "FFFFFF"


def set_run_font(run, latin="Aptos", east_asia="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, value, *, color=TEXT, bold=False, size=9.0, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(value))
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, txt, end])


def add_num_pages_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, txt, end])


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.16

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLUE)
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(8)

    for style_name, size, before, after in (
        ("Heading 1", 17, 10, 8),
        ("Heading 2", 13, 8, 5),
        ("Heading 3", 11, 6, 3),
    ):
        st = styles[style_name]
        st.font.name = "Aptos Display"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(TEXT)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Aptos"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(10)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.line_spacing = 1.12


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.25)

        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("PCB原理图自动解析项目综合技术报告")
        set_run_font(hr, size=8, color=MUTED)

        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("第 ")
        set_run_font(r, size=8, color=MUTED)
        add_page_field(fp)
        r = fp.add_run(" 页  共 ")
        set_run_font(r, size=8, color=MUTED)
        add_num_pages_field(fp)
        r = fp.add_run(" 页")
        set_run_font(r, size=8, color=MUTED)


def add_paragraph(doc, text="", *, bold_prefix=None, italic=False, color=TEXT, size=10.2, align=None, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.16
    if keep:
        p.paragraph_format.keep_with_next = True
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=size, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color)
        r.italic = italic
    return p


def add_bullets(doc, items, *, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if level:
            p.paragraph_format.left_indent = Inches(0.25 * level)
        if isinstance(item, tuple):
            lead, body = item
            r = p.add_run(lead)
            set_run_font(r, size=10, bold=True, color=TEXT)
            r = p.add_run(body)
            set_run_font(r, size=10, color=TEXT)
        else:
            r = p.add_run(item)
            set_run_font(r, size=10, color=TEXT)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r, size=10, color=TEXT)


def add_table(doc, headers, rows, widths=None, font_size=8.7, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        set_cell_text(cell, value, color=WHITE, bold=True, size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            cell.width = Inches(widths[idx])
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        set_keep_together(table.rows[-1])
        fill = PALEST_BLUE if r_idx % 2 else WHITE
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            set_cell_shading(cell, fill)
            set_cell_border(cell)
            set_cell_margins(cell)
            text_color = TEXT
            bold = False
            if status_col is not None and c_idx == status_col:
                bold = True
                val = str(value)
                if any(k in val for k in ("失败", "未通过", "缺失", "阻塞", "0 / 50")):
                    text_color = RED
                elif any(k in val for k in ("通过", "已验证", "可用", "权威")):
                    text_color = GREEN
                else:
                    text_color = AMBER
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or c_idx == status_col else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, value, color=text_color, bold=bold, size=font_size, align=align)
            if widths:
                cell.width = Inches(widths[c_idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def section_title(doc, number, title, subtitle=None, page_break=True):
    if page_break:
        doc.add_page_break()
    h = doc.add_paragraph(style="Heading 1")
    r = h.add_run(f"{number}  {title}")
    set_run_font(r, size=17, bold=True, color=TEXT)
    if subtitle:
        add_paragraph(doc, subtitle, color=MUTED, size=9.2, italic=True)


def subsection(doc, number, title):
    h = doc.add_paragraph(style="Heading 2")
    r = h.add_run(f"{number}  {title}")
    set_run_font(r, size=13, bold=True, color=TEXT)


def add_source_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("证据说明  ")
    set_run_font(r, size=8.5, bold=True, color=BLUE)
    r = p.add_run(text)
    set_run_font(r, size=8.5, color=MUTED)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    props = doc.core_properties
    props.title = "PCB原理图自动解析项目综合技术报告"
    props.subject = "赛题要求 技术架构 实验结果 当前状态与交付计划"
    props.author = "EDA挑战题目六项目组"
    props.keywords = "PCB 原理图 自动解析 Circuit-IR EDA"

    # Title page
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(p)
    r = p.add_run("PCB原理图自动解析项目综合技术报告")
    set_run_font(r, latin="Aptos Display", size=27, bold=True, color=BLUE)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("赛题要求  技术架构  实验结果  当前状态与交付计划")
    set_run_font(r, size=14, color=MUTED)
    doc.add_paragraph()
    meta = add_table(
        doc,
        ["项目", "内容"],
        [
            ("报告版本", "V1.0  整合清理版"),
            ("证据截止", "2026年9月16日"),
            ("覆盖范围", r"E:\EDA挑战题目六（含 archive\legacy-training）"),
            ("编制原则", "保留原始报告；以当前代码、配置、模型锁、审计产物和复现实测为准"),
        ],
        widths=[1.35, 5.65],
        font_size=9.3,
    )
    doc.add_paragraph()
    add_paragraph(
        doc,
        "本报告不再使用单一“完成度百分比”描述项目状态，而采用可复现证据、发布门禁和待验证项分层表述。",
        color=BLUE,
        size=10.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_page_break()

    # Contents
    h = doc.add_paragraph(style="Heading 1")
    r = h.add_run("目录")
    set_run_font(r, size=17, bold=True, color=TEXT)
    contents = [
        "1  执行摘要",
        "2  赛题要求与验收口径",
        "3  统一技术方案与系统架构",
        "4  数据资产与证据治理",
        "5  工程实现状态",
        "6  模型与实验结果",
        "7  验证结果与部署状态",
        "8  当前阻塞与推进优先级",
        "9  过时描述清理结果",
        "10  报告来源与处置说明",
        "11  结论与交付条件",
    ]
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(item)
        set_run_font(r, size=11, color=BLUE)
    add_source_note(doc, "正文中的“当前”均指证据截止日。历史文档只作为需求、设计或过程证据，不自动代表现行状态。")

    # 1 Executive summary
    section_title(doc, "1", "执行摘要", "以发布可用性为主线，对两处目录中的重复报告、状态报告与近期审计结果进行统一。")
    add_paragraph(
        doc,
        "项目已经具备从原理图图像到统一 Circuit-IR、图构建、规则校验和官方 JSON 序列化的完整工程链路。当前生产 Provider 能通过严格健康检查，四个任务模型已进入模型锁，单元与集成回归未发现失败项。",
    )
    add_paragraph(
        doc,
        "但“工程链路可运行”不等于“正式方案可发布”。最新严格 50 例发布运行仍因元件关键字段缺失和网络连通错误无法生成合规提交；本地 raw-v1 探针虽然给出可分析结果，但其数据集已被反复调参使用，不能当作独立验收集。正式发布状态因此仍为未通过。",
    )
    add_table(
        doc,
        ["判断维度", "当前结论", "关键证据"],
        [
            ("工程回归", "通过", "577 项测试通过，68 项因当前测试环境缺少 PyTorch 等依赖而跳过，0 失败"),
            ("运行时就绪", "通过", "release-candidate 配置下普通与 --require-ready 健康检查均返回 0"),
            ("模型锁", "已建立", "Component、OCR、Pin、Wire 四任务已登记；Junction 当前采用确定性矢量化"),
            ("严格发布门禁", "未通过", "50/50 样例均未通过 Validator，未形成官方 JSON"),
            ("端到端精度", "尚未正式确认", "raw-v1 本地代理分约 0.4815；数据暴露且非官方评测"),
            ("最新 Component 候选", "待独立确认", "验证集 F1 0.8410，较基线 +0.0200，FP 减少 57；尚未晋级模型锁"),
            ("目标环境部署", "部分验证", "Dockerfile 和契约测试存在；Linux/GPU 离线镜像构建与实机演练尚未完成"),
        ],
        widths=[1.3, 1.35, 4.35],
        status_col=1,
    )
    add_bullets(
        doc,
        [
            ("首要问题  ", "图拓扑和字段绑定仍决定最终得分，不能以单任务验证指标替代端到端验收。"),
            ("最新进展  ", "YOLO Component 边界精修在冻结验证集上显著降误检，具备进入独立确认的价值。"),
            ("发布原则  ", "只有独立数据、严格 Validator、官方或等价评分器和目标环境演练同时满足，才可改写为“可发布”。"),
        ],
    )

    # 2 Requirements
    section_title(doc, "2", "赛题要求与验收口径", "以下口径来自两处目录中内容一致的官方赛题指南 PDF。")
    subsection(doc, "2.1", "任务输入与输出")
    add_paragraph(
        doc,
        "系统输入为 PCB 原理图图像，输出为 JSON，核心对象包括 components、pins 和 nets。坐标原点位于左下角，x 轴向右、y 轴向上；重复位号需要追加 _1、_2 等后缀。网络中的每条 edge 必须恰含两个点，官方说明允许使用零长度边表示点或飞线端点。",
    )
    add_table(
        doc,
        ["对象", "核心匹配条件", "关键容差或规则"],
        [
            ("Component", "key、type、bbox 及文本属性", "bbox 容差 20 px；IC 比较 Name，电阻/电容比较 value"),
            ("Pin", "先按元件 key 对齐，再匹配引脚号", "名称严格；引脚点容差 5 px"),
            ("NetHypergraph", "把网络视为引脚集合进行全局匹配", "使用匈牙利匹配，强调连接关系而非仅线段外观"),
            ("NetLine", "只在已匹配网络内部比较线段", "端点容差 5 px；每条边必须有两个端点"),
            ("Flywire", "相同非 VCC/GND 标签可形成飞线关系", "VCC/GND 不按普通飞线处理"),
        ],
        widths=[1.2, 3.25, 2.55],
    )
    subsection(doc, "2.2", "评分与运行约束")
    add_table(
        doc,
        ["指标", "权重", "项目含义"],
        [
            ("Component F1", "30%", "元件识别与关键属性匹配"),
            ("Pin F1", "25%", "引脚位置、编号与名称匹配"),
            ("NetHypergraph F1", "35%", "网络拓扑，是精度部分权重最高的单项"),
            ("NetLine F1", "10%", "匹配网络内的线段几何质量"),
        ],
        widths=[1.8, 1.0, 4.2],
    )
    add_paragraph(doc, "精度部分占总成绩 95%，Token 指标占 5%。测试规模约 30 至 50 例，公开与黑盒比例约 1∶2，整体运行时间要求小于 2 小时。")
    add_source_note(
        doc,
        "两份官方指南的 SHA-256 均为 45EEDBC195AD71BBB03044C4606EA85CCF3ECCEBBB3D9B3304D6E325E1ACC5E8，内容可视为同一权威来源。",
    )

    # 3 Architecture
    section_title(doc, "3", "统一技术方案与系统架构", "将早期框架说明书中的稳定设计，与现行仓库中已经实现并验证的能力合并。")
    subsection(doc, "3.1", "端到端处理链")
    add_numbered(
        doc,
        [
            "输入治理与坐标变换  读取图像、建立变换链和底层坐标约定，避免模型坐标与官方坐标混用。",
            "多任务感知  分别执行 Component、OCR、Pin、Wire，并为每类结果保留置信度、来源和可追溯标识。",
            "实体关联  将文本、引脚和几何候选绑定到具体元件，形成统一 Circuit-IR。",
            "图构建  从线段、端点、交点、标签和电源符号构造物理图与电气图。",
            "规则校验  检查关键字段、悬空关系、连通一致性、重复位号和官方 JSON 契约。",
            "序列化与服务  通过 CLI 或 HTTP 输出严格 JSON，并提供模型就绪性、运行时后端和错误信息。",
        ],
    )
    subsection(doc, "3.2", "核心设计原则")
    add_table(
        doc,
        ["原则", "现行实现", "发布价值"],
        [
            ("统一中间表示", "Circuit-IR 串联感知、关联、图构建、校验和序列化", "减少模块间隐式约定"),
            ("物理图与电气图分层", "几何相交与逻辑同网分别建模", "更贴近 NetHypergraph 评分"),
            ("失败关闭", "显式 CUDA 请求发生 CPU 回退时判为失败；严格模式拒绝不合规输出", "避免静默降级或伪成功"),
            ("模型锁与哈希绑定", "模型注册、锁文件和后端就绪性共同决定可用任务", "让结果可复现、可审计"),
            ("证据分层", "单任务验证、raw-v1 探针、严格发布运行、官方评测分别标注", "防止将局部指标误写成最终成绩"),
        ],
        widths=[1.45, 3.25, 2.3],
    )
    subsection(doc, "3.3", "历史设计的保留与调整")
    add_paragraph(
        doc,
        "早期两份 Word 说明书对 L0 至 L7 分层、OCR、图推理、校验器、序列化和容器化提出了完整方案，这些内容继续作为架构背景。其“两周/十周计划”、推荐模型、人员安排和当时的完成度不再作为现状结论；Qwen 也由早期“待实现”调整为“代码已具备、发布配置关闭、缺少收益证据”。",
    )

    # 4 Data
    section_title(doc, "4", "数据资产与证据治理", "数据重叠和反复调参会直接改变实验结论的可信级别。")
    subsection(doc, "4.1", "Dataset300 审计")
    add_table(
        doc,
        ["检查项", "结果", "解释"],
        [
            ("图像与标注", "300 + 300，哈希核验完成", "基本资产齐备"),
            ("严格审计", "19 项警告", "8 项扩展名/格式不一致，11 组重复或近重复"),
            ("分组划分", "240 / 30 / 30", "训练、验证、测试按重复组隔离"),
            ("派生任务记录", "147,758 条", "Component 12,789；Pin 29,648；OCR 48,464；Wire 14,895；Junction 41,962"),
        ],
        widths=[1.35, 1.8, 3.85],
    )
    subsection(doc, "4.2", "公开数据与 Dataset300 的关系")
    add_bullets(
        doc,
        [
            "公开数据的前 200 张图像与 Dataset300 前 200 张逐一哈希完全相同，不能按两套独立数据直接拼接统计。",
            "新增 100 张均来自 datasheet 域，按重复组计算只新增约 90 个独立组；总有效独立组约 287。",
            "公开 10 张 GT 图像也存在于公开训练集，且目标 JSON 表示并非逐字节相同，不适合作为独立测试集。",
            "competition JSON、CVJsonStd、EDS 与 SPICE 是不同表示层；比较时必须统一对象定义与坐标语义。",
        ],
    )
    subsection(doc, "4.3", "release50 暴露结论")
    add_paragraph(
        doc,
        "现有 release50 由 30 例 test 与 20 例 val 组成，已经被多轮优化和阈值选择使用。它仍适合做回归对比和故障定位，但不再具备独立留出集资格。报告中的 release50 指标统一标为“本地代理”或“回归证据”，不作为比赛成绩或最终泛化结论。",
    )
    add_source_note(doc, "数据结论来自 Dataset300 审计、数据集对比汇总和 split exposure audit；截止 2026年9月12日。")

    # 5 Implementation
    section_title(doc, "5", "工程实现状态", "状态按“已验证、已实现待发布验证、实验性或缺失”区分。")
    add_table(
        doc,
        ["模块", "当前状态", "边界说明"],
        [
            ("数据转换与任务记录", "已验证", "Dataset300 审计、分组划分及五任务记录已生成"),
            ("Circuit-IR 与坐标变换", "已验证", "关键契约有单元和集成回归覆盖"),
            ("实体关联", "已验证", "支持元件、文本和引脚绑定，但端到端关键字段仍偏弱"),
            ("图构建与校验器", "已验证", "共用共线容差与间隙契约；严格 release50 仍暴露连通错误"),
            ("官方 JSON 序列化", "已验证", "重复位号后缀、edge 两点等契约已覆盖"),
            ("Production Provider", "运行时可用", "四任务模型注册并可加载，健康检查通过"),
            ("模型锁", "已建立", "Component、OCR、Pin、Wire 已锁；最新 YOLO 候选未晋级"),
            ("Junction", "确定性方案可用", "学习式 Junction 因真实验证集类别支持不足未推广"),
            ("CLI 与 HTTP API", "已验证", "基本调用、错误处理与健康检查路径具备"),
            ("Docker", "部分验证", "Dockerfile 和契约测试存在；尚缺目标 Linux/GPU 离线演练"),
            ("Qwen 编排", "实验性关闭", "代码已存在，但 release-candidate 配置关闭，尚无正向 A/B 证据"),
            ("官方评测器", "缺失", "本地未找到官方 evaluator 或等价最终评分环境"),
        ],
        widths=[1.45, 1.45, 4.1],
        status_col=1,
    )
    add_paragraph(
        doc,
        "现行 release-candidate 配置使用 Component 512 像素切片和 64 像素重叠，Component 阈值 0.3，OCR 检测阈值 0.5，Pin 阈值 0.75，Wire 阈值 0.6，并在严格校验模式下要求零错误。上述参数属于当前运行配置，不等同于所有参数都已在独立集上定型。",
    )

    # 6 Metrics
    section_title(doc, "6", "模型与实验结果", "任务级模型指标、端到端代理指标和实验候选必须分别解读。")
    subsection(doc, "6.1", "当前模型锁的任务级指标")
    add_table(
        doc,
        ["任务", "锁定指标", "解读"],
        [
            ("Component", "mAP50 0.2501；P 0.0393；R 0.3166", "当前生产锁仍为旧候选，明显落后于近期 YOLO 实验"),
            ("OCR", "归一化精确率 0.5307；CER 0.3006；检测 Hmean 0.3669", "单任务已脱离早期 CER=1.0 状态，端到端字段绑定仍弱"),
            ("Pin", "F1@4px 0.6376；P 0.8508；R 0.5098", "精度较高、召回仍是约束"),
            ("Wire", "骨架 F1 0.9402；连通性 0.5789", "像素/骨架质量较好，不能替代电气图正确性"),
        ],
        widths=[1.15, 2.75, 3.1],
    )
    add_source_note(doc, "以上数值来自 models/models.lock.json 对应的已祝福模型，不是官方端到端得分。")

    subsection(doc, "6.2", "raw-v1 本地端到端探针")
    add_table(
        doc,
        ["指标", "数值", "限制"],
        [
            ("Component F1", "0.7807", "release50 已暴露，非独立"),
            ("Pin F1", "0.5166", "字段绑定和召回仍影响结果"),
            ("Graph pairwise F1", "0.2013", "仅为本地近似，不等价官方 NetHypergraph"),
            ("Wire F1", "0.4768", "与任务级骨架指标口径不同"),
            ("加权代理分", "约 0.4815", "按官方四项权重近似计算，不含 Token，不是官方成绩"),
            ("运行时间", "1620.44 秒 / 50 例", "满足 2 小时上限，但属于优化探针路径"),
        ],
        widths=[1.65, 1.55, 3.8],
    )
    add_paragraph(
        doc,
        "严格字段命中仍很低：component designator 0.1769、value 0.1616、pin number 0.0544、pin name 0.0063。由此可见，当前主要差距不仅是框线检测，还包括 OCR、字段归一化与对象绑定。",
    )

    doc.add_page_break()
    subsection(doc, "6.3", "最新 YOLO Component 候选")
    add_table(
        doc,
        ["阶段", "关键结果", "当前结论"],
        [
            ("YOLO11n v1 训练", "最佳 epoch 92：P 0.8247，R 0.7090，mAP50 0.8176", "训练验证指标优秀，但与真实切片回放口径不同"),
            ("YOLO11n v2 精度版", "最佳 epoch 7：P 0.8104，R 0.7377，mAP50 0.8166", "候选，不在生产锁"),
            ("切片校准", "阈值 0.629552：F1 0.8269，P 0.8027，R 0.8525", "同一冻结验证集上选择阈值，存在调参偏差"),
            ("边界精修 A/B", "F1 0.8210 → 0.8410；P 0.7658 → 0.8013；FP 301 → 244", "减少 57 个 FP，召回 0.8849 不变"),
            ("推广状态", "validation_candidate_not_promoted", "缺独立集确认、下游 Pin/Graph 回放与正式模型锁更新"),
        ],
        widths=[1.45, 3.45, 2.1],
        status_col=2,
    )
    add_paragraph(
        doc,
        "该候选是当前最有价值的近期进展，但只能表述为“验证集上的晋级候选”。下一步应冻结配置，在未参与阈值选择的独立数据上复测，并完整回放 Pin、Graph、Validator 和官方 JSON。",
    )

    # 7 Validation deployment
    section_title(doc, "7", "验证结果与部署状态", "本节记录在证据截止日重新执行的回归与健康检查。")
    subsection(doc, "7.1", "自动化回归")
    add_table(
        doc,
        ["测试层级", "结果", "说明"],
        [
            ("Unit", "524 通过", "68 项跳过，主要因 .venv-test 未安装 PyTorch 等重依赖"),
            ("Integration", "53 通过", "核心流水线、契约与接口回归未失败"),
            ("合计", "577 通过  68 跳过  0 失败", "跳过项不应被计作通过，发布前需在完整环境补跑"),
        ],
        widths=[1.4, 2.35, 3.25],
    )
    subsection(doc, "7.2", "生产健康检查")
    add_paragraph(
        doc,
        "在 .venv-train、pipeline.release-candidate.yaml 和 models.lock.json 下，普通健康检查与 --require-ready 严格健康检查均返回 0。Component、OCR、Pin、Wire 的注册、文件校验与运行时 smoke test 均达到当前 ready 口径。",
    )
    add_paragraph(
        doc,
        "GPU 运行仍有环境兼容风险：当前 ONNX Runtime 1.28 尝试加载 CUDA 13 相关 cublasLt64_13.dll，而主机 CUDA 为 12.6。auto 模式会回退 CPU 并保持就绪；显式要求 CUDA 时已改为失败关闭。PyTorch CUDA 可用不代表 ONNX Runtime GPU 后端已经兼容。",
    )
    subsection(doc, "7.3", "严格发布运行")
    add_table(
        doc,
        ["检查项", "结果", "发布含义"],
        [
            ("50 例序列化", "0 / 50 成功", "严格 Validator 拒绝全部输出"),
            ("元件关键字段", "约 1,699 项缺失", "位号、数值或名称仍需提升"),
            ("断连线段", "约 1,681 项", "网络端点吸附、交点和拓扑重建仍是主阻塞"),
            ("总耗时", "891.39 秒", "小于 2 小时，但耗时合格不能替代输出合规"),
            ("P95 单例", "53.04 秒", "长尾仍需在目标设备复核"),
            ("发布硬门禁", "未通过", "registry 中 production_ready 仍为 false"),
        ],
        widths=[1.55, 1.65, 3.8],
        status_col=1,
    )
    add_source_note(doc, "严格 50 例证据来自 2026年8月31日发布运行；健康检查和回归测试于 2026年9月16日重新执行。")

    # 8 blockers
    section_title(doc, "8", "当前阻塞与推进优先级", "按对官方总分、发布合规和证据可信度的影响排序。")
    add_table(
        doc,
        ["优先级", "工作项", "验收条件"],
        [
            ("P0", "修复严格 Validator 的关键字段缺失", "50 例严格运行可稳定输出；不得以放宽关键字段规则换取通过"),
            ("P0", "提升网络图连通性", "断连、错误合并和标签飞线问题显著下降；Graph 指标在独立集确认"),
            ("P0", "建立独立端到端确认集", "样本未参与训练、阈值选择和多轮调参；重复组隔离可审计"),
            ("P1", "推进 YOLO Component 候选", "独立集保留 F1 与误检改善，并完成下游 Pin/Graph 回放后更新模型锁"),
            ("P1", "改善 OCR 与字段绑定", "designator、value、pin number、pin name 严格命中率按对象类别提升"),
            ("P1", "补齐官方或等价评分器", "评分口径、坐标、匈牙利匹配和异常输出与官方说明一致"),
            ("P2", "完成目标环境演练", "Linux/GPU Docker 离线构建、模型加载、50 例运行和日志归档完成"),
            ("P2", "评估 Qwen 编排", "只有在固定数据和相同门禁下证明净收益后才启用"),
        ],
        widths=[0.75, 2.35, 3.9],
    )
    add_paragraph(
        doc,
        "计划不再以“还需几天”或“总体完成度”衡量，而以可交付物和硬门禁衡量。任何新模型或规则必须同时记录代码版本、配置、模型哈希、数据切分、运行环境和端到端结果。",
    )

    # 9 stale cleanup
    section_title(doc, "9", "过时描述清理结果", "下表列出被删除或改写的主要口径；原始文件仍保留作历史记录。")
    add_table(
        doc,
        ["历史描述", "清理原因", "统一后的现行表述"],
        [
            ("Production Provider 尚未实现", "与当前健康检查矛盾", "Provider 已实现并通过 ready 检查；正式发布仍未过门禁"),
            ("models.lock.json 未更新", "四任务模型现已进入模型锁", "锁已建立；最新 YOLO Component 候选尚未晋级"),
            ("Dockerfile 尚未开始", "仓库已有 Dockerfile 与契约测试", "容器骨架存在；目标 Linux/GPU 离线演练待补"),
            ("Qwen 框架未实现", "现已有编排代码", "代码存在但 release 配置关闭，尚无收益 A/B"),
            ("Component mAP 约 0.05 至 0.08", "属于早期失败训练", "生产锁 mAP50 0.2501；近期 YOLO 候选验证指标更高但未晋级"),
            ("Pin 无预测或 PCK 为 0", "已被现行模型结果替代", "锁定 Pin F1@4px 0.6376；召回仍需提升"),
            ("Wire F1 约 0.18 至 0.30", "属于旧实验口径", "锁定骨架 F1 0.9402，但端到端 Wire 代理约 0.4768"),
            ("OCR CER 为 1.0", "属于早期训练状态", "锁定 CER 0.3006；端到端文本字段仍是瓶颈"),
            ("所有非训练模块均无问题", "绝对化且与严格发布失败矛盾", "577 项回归通过，但 68 项跳过，严格发布仍失败"),
            ("流水线 10 图仅需 0.1 秒", "多为 golden/debug/启发式路径", "生产类运行需按完整 Provider 和 50 例实测报告"),
            ("总体完成度 55%、65% 等", "定义冲突且不可审计", "改为工程就绪、模型锁、端到端、部署和发布门禁矩阵"),
            ("约 4 天即可完成", "缺少依赖与验收前提", "改为 P0/P1/P2 及明确验收条件"),
            ("关闭 Junction 缺类门禁", "会掩盖数据支持不足", "保留严格门禁，当前使用确定性矢量化，不强行推广学习模型"),
        ],
        widths=[2.05, 1.95, 3.0],
        font_size=8.2,
    )

    # 10 sources
    section_title(doc, "10", "报告来源与处置说明", "对两个目录中的报告按权威性、时效性和用途进行归档解释。")
    subsection(doc, "10.1", "权威与现行证据")
    add_table(
        doc,
        ["来源", "处置", "用途"],
        [
            ("两份官方赛题指南 PDF", "权威来源", "任务、JSON、评分、容差和运行约束"),
            (r"E:\EDA挑战题目六\models\models.lock.json", "现行证据", "锁定模型、哈希与任务指标"),
            (r"E:\EDA挑战题目六\models\registry.yaml", "现行证据", "runtime_ready 与 production_ready 状态"),
            (r"E:\EDA挑战题目六\configs\pipeline.release-candidate.yaml", "现行证据", "当前发布候选配置"),
            ("optimization-20260908 与近期 artifacts/work 审计", "现行过程证据", "契约修复、数据暴露、raw-v1 与最新候选实验"),
        ],
        widths=[3.1, 1.25, 2.65],
        status_col=1,
    )
    subsection(doc, "10.2", "历史报告归并")
    add_table(
        doc,
        ["文件或文件组", "归并方式"],
        [
            ("PCB原理图自动解析_系统框架详细说明书.docx", "保留稳定架构与人员/职责背景；移除当时的状态推断"),
            ("赛题六_PCB原理图自动解析_落地部署实施手册.docx", "保留部署设计与分层方案；计划周期改为验收门禁"),
            ("赛题分析报告.md", "保留赛题拆解；模型选型和当前状态按现行证据重写"),
            ("当前状态与下一步计划.md、完成度报告.md、系统构建状态报告.md", "合并为工程状态矩阵，删除互相冲突的完成度百分比"),
            ("系统模块完成度.md、系统完成度分析.md、系统完成度重新分析.md", "作为历史快照归档，不再用于发布判断"),
            ("修复与训练报告.md、训练优化结果.md、组件完成度报告.md", "旧模型指标标为历史，使用模型锁和最新候选实验替代"),
            ("README.md、docs/implementation_status.md、系统完成度分析与下一步方案.md", "提取仍有效的模块说明；时间点较早的状态不直接沿用"),
            ("架构阻塞与硬件优化审查、PCBnet 论文驱动建议", "保留为技术建议；只把有当前产物或实测支撑的内容写入现状"),
            ("部署与训练手册_PCBnet增强版_2026-09-09.md", "保留操作流程；实验结果以锁文件与最新审计为准"),
            ("optimization-20260908 下的 7 份审计报告", "作为 9 月工程修复和证据治理的主要过程来源"),
        ],
        widths=[3.65, 3.35],
        font_size=8.4,
    )
    add_source_note(doc, "两处目录中的原始报告均未删除或覆盖。本报告是统一入口；历史文档继续承担过程追溯功能。")

    # 11 conclusion
    section_title(doc, "11", "结论与交付条件", "以可复现证据定义下一次状态升级。")
    add_paragraph(
        doc,
        "截至 2026年9月16日，项目已经完成主要工程骨架、运行时装配和大部分契约治理，具备继续集中优化的基础。项目当前最准确的定位是“工程运行时就绪、正式发布未就绪”。",
    )
    add_paragraph(
        doc,
        "下一阶段不应再扩大功能面，而应围绕三个闭环推进：一是用独立数据确认 YOLO Component 候选；二是解决严格 Validator 暴露的字段缺失与图连通问题；三是在官方或等价评分器和目标 Linux/GPU 环境中完成可复现演练。",
    )
    add_table(
        doc,
        ["状态升级", "必须满足的条件"],
        [
            ("候选模型可晋级", "独立集确认提升；下游 Pin/Graph 不退化；配置与权重哈希入锁"),
            ("端到端可验收", "严格 JSON 通过率达到目标；Component、Pin、NetHypergraph、NetLine 按统一口径复现"),
            ("部署可交付", "目标 Linux/GPU 容器离线构建和全量运行成功；显式 CUDA 不发生静默回退"),
            ("比赛可提交", "官方 evaluator 或经核对的等价实现可运行；30 至 50 例总耗时小于 2 小时；报告与产物版本一致"),
        ],
        widths=[1.65, 5.35],
    )
    add_paragraph(
        doc,
        "本报告已将两个目录中的需求、方案、实验、工程状态和近期优化统一到同一证据口径。后续更新时，应优先更新本报告中的状态矩阵和证据表，而不是继续新增彼此独立、口径不一致的完成度报告。",
        color=BLUE,
        size=10.5,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
